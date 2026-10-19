"""Write the daily operations report as an Excel workbook."""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill


def _sheet(wb, title, header, rows):
    ws = wb.create_sheet(title)
    ws.append(header)
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="DDDDDD")
    for r in rows:
        ws.append(r)
    for i, h in enumerate(header, 1):
        ws.column_dimensions[chr(64 + i)].width = max(14, len(h) + 4)
    return ws


def write(path, rec, money, aml_alerts):
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    ws.append(["Measure", "Value"])
    ws["A1"].font = ws["B1"].font = Font(bold=True)
    for k, v in [("internal trades", rec["internal"]), ("broker rows", rec["broker_rows"]), ("matched", rec["matched"]),
                 ("breaks", len(rec["breaks"])), ("client money status", money["status"]), ("client money gap (ore)", money["gap_ore"]),
                 ("AML alerts", len(aml_alerts))]:
        ws.append([k, v])
    _sheet(wb, "Breaks", ["trade_id", "kind", "detail"], [[b["trade_id"], b["kind"], b["detail"]] for b in rec["breaks"]])
    _sheet(wb, "ClientMoney", ["owed_ore", "bank_ore", "gap_ore", "status"], [[money["owed_ore"], money["bank_ore"], money["gap_ore"], money["status"]]])
    _sheet(wb, "AML", ["client", "rule", "reason"], [[a["client"], a["rule"], a["reason"]] for a in aml_alerts])
    wb.save(path)

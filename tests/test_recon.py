from collections import Counter

from recon import aml, client_money, gen, match, report


def test_every_planted_break_is_found_with_the_right_kind():
    internal, broker, planted = gen.make_day()
    found = Counter(b["kind"] for b in match.reconcile(internal, broker)["breaks"])
    assert dict(found) == planted


def test_matched_plus_trades_with_a_break_equals_internal_trades():
    internal, broker, _ = gen.make_day()
    rec = match.reconcile(internal, broker)
    broken_internal = {b["trade_id"] for b in rec["breaks"] if b["kind"] in ("missing_at_broker", "quantity_mismatch", "price_mismatch")}
    assert rec["matched"] + len(broken_internal) == rec["internal"]


def test_price_difference_inside_tolerance_is_not_a_break():
    t = {"trade_id": "T1", "qty": 5, "price_ore": 10_000}
    b = {"trade_id": "T1", "qty": 5, "price_ore": 10_050}
    assert match.reconcile([t], [b])["breaks"] == []


def test_client_money_shortfall_is_flagged():
    assert client_money.check({"C1": 100, "C2": 200}, 250)["status"] == "SHORTFALL"
    assert client_money.check({"C1": 100, "C2": 200}, 300)["status"] == "ok"


def test_client_money_in_the_demo_data_is_covered():
    assert client_money.check(*gen.make_client_money())["status"] == "ok"


def _rows():
    return gen.make_cash_movements()


def test_structuring_pattern_is_caught_for_the_planted_client():
    hits = [a for a in aml.alerts(_rows()) if a["rule"] == "STRUCTURING"]
    assert [h["client"] for h in hits] == ["C900"]


def test_pass_through_large_and_high_risk_country_rules():
    rules = {(a["client"], a["rule"]) for a in aml.alerts(_rows())}
    assert ("C901", "PASS_THROUGH") in rules
    assert ("C901", "LARGE_MOVEMENT") in rules
    assert ("C902", "HIGH_RISK_COUNTRY") in rules


def test_ordinary_clients_raise_no_alert():
    ordinary = [r for r in _rows() if r["client"] not in ("C900", "C901", "C902")]
    assert [a for a in aml.alerts(ordinary) if a["rule"] != "LARGE_MOVEMENT"] == []


def test_report_workbook_has_the_four_sheets(tmp_path):
    from openpyxl import load_workbook
    internal, broker, _ = gen.make_day()
    rec = match.reconcile(internal, broker)
    p = tmp_path / "r.xlsx"
    report.write(p, rec, client_money.check(*gen.make_client_money()), aml.alerts(_rows()))
    wb = load_workbook(p)
    assert wb.sheetnames == ["Summary", "Breaks", "ClientMoney", "AML"]
    assert wb["Breaks"].max_row == len(rec["breaks"]) + 1
